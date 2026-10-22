"""Write edits into a sheet, changing nothing else.

openpyxl cannot be used to save: it drops every formula result Excel
stored, and NOMAD's parser reads only those (PR 3). So only the target
cells' XML is rewritten; every other byte of every part stays as it was.

Formulas that read an edited cell get their stored result recomputed
(formulas.py). Before a result is trusted, the calculator must reproduce
Excel's stored result from the original inputs; otherwise the edit is
refused. Every write is verified before it is returned.
"""

import datetime
import io
import math
import re
import zipfile
from dataclasses import dataclass
from xml.etree import ElementTree
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from sand.sheet_edit import formulas
from sand.sheet_edit.formulas import ExcelError, Ref, Unsupported
from sand.sheet_edit.workbook import (
    FIRST_SAMPLE_ROW,
    HEADER_ROW,
    CellChange,
    changes_between,
)

NewValue = str | int | float | bool | None


@dataclass(frozen=True)
class Write:
    row: int  # 1-based
    column: int  # 1-based
    value: NewValue  # converted to the original cell's type; None clears


class WriteRefused(Exception):
    """The edit cannot be written safely; nothing was written."""


class VerifyFailed(Exception):
    """The written file differs from what was intended; it must not be used."""


@dataclass(frozen=True)
class WriteResult:
    xlsx: bytes
    # every cell that changed, recomputed formula results included
    changes: list[CellChange]


def write_cells(xlsx: bytes, writes: list[Write]) -> WriteResult:
    book = load_workbook(io.BytesIO(xlsx))
    values = load_workbook(io.BytesIO(xlsx), data_only=True)
    ws, stored = book.worksheets[0], values.worksheets[0]

    new: dict[tuple[int, int], NewValue] = {}
    for w in writes:
        key = (w.row, w.column)
        # checked before ws.cell, which would create (and later scan) the cell
        if not (1 <= w.row <= ws.max_row and 1 <= w.column <= ws.max_column):
            raise WriteRefused(f'row {w.row}, column {w.column} is outside the sheet')
        if key in new:
            raise WriteRefused(f'{_coord(*key)} is written twice')
        if w.row < FIRST_SAMPLE_ROW:
            raise WriteRefused(f'{_coord(*key)} is a header; headers are not edited')
        cell = ws.cell(w.row, w.column)
        if cell.data_type == 'f':
            raise WriteRefused(f'{cell.coordinate} is calculated by a formula')
        value = _typed(w.value, _kind(ws, w.row, w.column), cell.coordinate)
        if value != cell.value or type(value) is not type(cell.value):
            new[key] = value

    results = _recompute(book, stored, new)

    changes = [
        *_value_changes(ws, new),
        *(
            # an empty text result is stored as <v></v>, which reads back as None
            CellChange(
                ws.title, r, c, 'computed value', stored.cell(r, c).value, v or None
            )
            for (r, c), v in results.items()
        ),
    ]
    if not changes:
        return WriteResult(xlsx, [])

    after = _patch(xlsx, new, results)
    _verify(xlsx, after, changes, before_books=(book, values))
    return WriteResult(after, changes)


def _coord(row: int, column: int) -> str:
    return f'{get_column_letter(column)}{row}'


# --- types -------------------------------------------------------------------


def _kind(ws, row: int, column: int) -> str:
    """'number' | 'text' | 'bool' | 'date' | 'any', from the cell, else its
    column."""
    value = ws.cell(row, column).value
    values = (
        [value]
        if value is not None
        else [
            ws.cell(r, column).value
            for r in range(FIRST_SAMPLE_ROW, ws.max_row + 1)
            if ws.cell(r, column).value is not None
        ]
    )
    kinds = {
        'bool'
        if isinstance(v, bool)
        else 'number'
        if isinstance(v, int | float)
        else 'date'
        if isinstance(v, datetime.date | datetime.time | datetime.timedelta)
        else 'text'
        for v in values
    }
    return kinds.pop() if len(kinds) == 1 else 'any'


def _typed(value: NewValue, kind: str, coord: str) -> NewValue:  # noqa: PLR0911
    """`value` in the type the cell holds: a number stays a number, text
    stays text even when it looks like a number."""
    if value is None or value == '':
        return None
    if kind == 'date':
        # writing text would silently turn the date into a string
        raise WriteRefused(f'{coord} holds a date; sand does not edit dates yet')
    if kind == 'text':
        return formulas.as_text(value)
    if kind == 'bool':
        if isinstance(value, bool):
            return value
        if str(value).strip().upper() in ('TRUE', 'FALSE'):
            return str(value).strip().upper() == 'TRUE'
        raise WriteRefused(f'{coord} holds TRUE/FALSE; {value!r} is neither')
    try:
        number = _number(value)
    except ValueError:
        if kind == 'any':  # an empty column: keep what was said
            return value
        raise WriteRefused(
            f'{coord} holds numbers; {value!r} is not a number'
        ) from None
    return number


def _number(value: NewValue) -> int | float:
    if isinstance(value, bool):
        raise ValueError(value)
    f = float(value.strip()) if isinstance(value, str) else float(value)
    if not math.isfinite(f):
        raise ValueError(value)
    return int(f) if f.is_integer() else f


def _data_type(value: NewValue) -> str:
    """openpyxl's data_type for a written value."""
    if isinstance(value, bool):
        return 'b'
    if isinstance(value, str):
        return 's'
    return 'n'


def _value_changes(ws, new: dict[tuple[int, int], NewValue]) -> list[CellChange]:
    changes = []
    for (r, c), value in new.items():
        cell = ws.cell(r, c)
        if cell.value != value:
            changes.append(CellChange(ws.title, r, c, 'value', cell.value, value))
        if cell.data_type != _data_type(value):
            changes.append(
                CellChange(ws.title, r, c, 'type', cell.data_type, _data_type(value))
            )
    return changes


# --- formula results ---------------------------------------------------------


# (first row, first column, last row, last column); open ranges like A:A
# or 3:3 run to infinity
_Rect = tuple[float, float, float, float]
_EVERYTHING: list[_Rect] = [(1, 1, math.inf, math.inf)]
_INF = math.inf

_STRING_RE = re.compile(r'"(?:[^"]|"")*"')
_CELL_RANGE_RE = re.compile(
    r"(?:(?:'[^']+'|[A-Za-z0-9_.]+)!)?"  # optional sheet
    r'\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?'
)
_COLUMNS_RE = re.compile(
    r"(?:(?:'[^']+'|[A-Za-z0-9_.]+)!)?\$?([A-Z]{1,3}):\$?([A-Z]{1,3})\b"
)
_ROWS_RE = re.compile(r"(?:(?:'[^']+'|[A-Za-z0-9_.]+)!)?\$?(\d+):\$?(\d+)\b")
_FUNCTION_RE = re.compile(r'(?:_xlfn\.|_xlws\.)?[A-Za-z][A-Za-z0-9.]*\(')
# functions whose inputs are computed at run time, not written in the text
_DYNAMIC = {'INDIRECT', 'OFFSET'}


def _reads(text: str) -> list[_Rect]:
    """The cells a formula may read.

    Exact for formulas the calculator parses. For others, every cell and
    range written in the text, so an edit that might feed the formula is
    caught (and refused) rather than missed. When the text leaves inputs
    open (a defined name like `Speed`, INDIRECT, OFFSET, a table
    reference), the formula may read anything.
    """
    try:
        return [(r.row, r.column, r.row, r.column) for r in formulas.references(text)]
    except Unsupported:
        pass
    rest = _STRING_RE.sub(' ', text.lstrip('='))
    if any(
        name.rstrip('(').upper().split('.')[-1] in _DYNAMIC
        for name in _FUNCTION_RE.findall(rest)
    ):
        return _EVERYTHING
    rest = _FUNCTION_RE.sub(' ', rest)

    rects: list[_Rect] = []
    for m in _COLUMNS_RE.finditer(rest):
        c1, c2 = sorted(column_index_from_string(g) for g in m.groups())
        rects.append((1, c1, _INF, c2))
    rest = _COLUMNS_RE.sub(' ', rest)
    for m in _ROWS_RE.finditer(rest):
        r1, r2 = sorted(int(g) for g in m.groups())
        rects.append((r1, 1, r2, _INF))
    rest = _ROWS_RE.sub(' ', rest)
    for m in _CELL_RANGE_RE.finditer(rest):
        r1, c1 = int(m.group(2)), column_index_from_string(m.group(1))
        r2 = int(m.group(4)) if m.group(3) else r1
        c2 = column_index_from_string(m.group(3)) if m.group(3) else c1
        rects.append((min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)))
    rest = _CELL_RANGE_RE.sub(' ', rest)

    # a word or table reference left over is a name sand cannot resolve
    if re.search(r'[A-Za-z_\[]', rest.replace('TRUE', '').replace('FALSE', '')):
        return _EVERYTHING
    return rects


def _hits(rects: list[_Rect], cells: set[tuple[int, int]]) -> bool:
    return any(
        r1 <= r <= r2 and c1 <= c <= c2 for r1, c1, r2, c2 in rects for r, c in cells
    )


def _formula_reads(value) -> list[_Rect]:
    """An array formula's text still says what it reads; a data table's
    does not."""
    text = value if isinstance(value, str) else getattr(value, 'text', None)
    return _reads(text) if text else _EVERYTHING


def _formula_text(value) -> str | None:
    """A cell's formula text; None for array and data-table formulas."""
    return value if isinstance(value, str) else None


def _recompute(book, stored, new: dict[tuple[int, int], NewValue]) -> dict:
    """New stored results of every formula that reads an edited cell,
    directly or through another formula. Only results that change."""
    if not new:
        return {}
    ws = book.worksheets[0]
    _refuse_other_sheets_reading(book, new)

    raw = {
        (cell.row, cell.column): cell.value
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == 'f'
    }
    texts = {key: _formula_text(value) for key, value in raw.items()}
    reads = {key: _formula_reads(value) for key, value in raw.items()}

    affected: set[tuple[int, int]] = set()
    frontier = set(new)
    while frontier:
        frontier = {
            key for key in texts if key not in affected and _hits(reads[key], frontier)
        }
        affected |= frontier

    results: dict[tuple[int, int], formulas.Value] = {}
    in_progress: set[tuple[int, int]] = set()

    def name_of(key: tuple[int, int]) -> str:
        header = ws.cell(HEADER_ROW, key[1]).value
        return f'{header!r} ({_coord(*key)})' if header else _coord(*key)

    def original(ref: Ref):
        return stored.cell(ref.row, ref.column).value

    def current(ref: Ref):
        key = (ref.row, ref.column)
        if key in new:
            return new[key]
        if key in affected:
            return result(key)
        return original(ref)

    def result(key: tuple[int, int]):
        if key in results:
            return results[key]
        if key in in_progress:
            raise WriteRefused(f'{name_of(key)} is part of a circular calculation')
        if texts[key] is None:
            raise WriteRefused(
                f'{name_of(key)} is an array or table formula; sand cannot compute it'
            )
        in_progress.add(key)
        results[key] = _new_result(
            name_of(key), texts[key], stored.cell(*key).value, original, current
        )
        in_progress.discard(key)
        return results[key]

    for key in sorted(affected):
        result(key)
    return {
        key: value
        for key, value in results.items()
        if _result_differs(value, stored.cell(*key).value)
    }


def _result_differs(new: formulas.Value, old: formulas.Value) -> bool:
    if new in ('', None) and old in ('', None):
        return False
    return not formulas.same_result(new, old) or type(new) is not type(old)


def _refuse_other_sheets_reading(book, new: dict) -> None:
    """Formulas on other sheets that point at this one would go stale."""
    title = book.worksheets[0].title
    for other in book.worksheets[1:]:
        for row in other.iter_rows():
            for cell in row:
                text = _formula_text(cell.value) if cell.data_type == 'f' else None
                if cell.data_type == 'f' and (text is None or '!' in text):
                    raise WriteRefused(
                        f'a formula on sheet {other.title!r} ({cell.coordinate}) may '
                        f'read {title!r}; sand cannot keep it up to date'
                    )


def _new_result(name: str, text: str, excel, original, current):
    """The formula's new result, after proving the calculator reproduces
    `excel`, the result Excel stored, from the original inputs."""
    try:
        check = formulas.evaluate(text, original)
    except Unsupported as exc:
        raise WriteRefused(f'{name} is a formula sand cannot compute ({exc})') from None
    except ExcelError as exc:
        check = exc
    # Excel stores an empty text result as <v></v>, which reads as None:
    # only a missing result with something to show is "no stored result"
    if excel is None and check not in ('', None):
        raise WriteRefused(
            f'{name} has no stored result to check sand against; '
            'open and save the file in Excel first'
        )
    if isinstance(check, ExcelError) or not formulas.same_result(check, excel):
        raise WriteRefused(
            f'{name}: sand computes {check!r} but Excel stored {excel!r}; '
            'sand cannot update this formula'
        )
    try:
        value = formulas.evaluate(text, current)
    except ExcelError as exc:
        raise WriteRefused(f'{name} would become an Excel error: {exc}') from None
    except Unsupported as exc:
        raise WriteRefused(
            f'{name} cannot be computed with the new value ({exc})'
        ) from None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return value


# --- XML ---------------------------------------------------------------------

_NS = {
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'p': 'http://schemas.openxmlformats.org/package/2006/relationships',
}
_CELL_RE = re.compile(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', re.S)
_REF_RE = re.compile(r'\br="([A-Z]+)(\d+)"')


def _first_sheet_part(z: zipfile.ZipFile) -> str:
    workbook = ElementTree.fromstring(z.read('xl/workbook.xml'))
    sheet = workbook.find('m:sheets/m:sheet', _NS)
    rid = sheet.get(f'{{{_NS["r"]}}}id')
    rels = ElementTree.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    target = next(
        rel.get('Target')
        for rel in rels.findall('p:Relationship', _NS)
        if rel.get('Id') == rid
    )
    return target.lstrip('/') if target.startswith('/') else f'xl/{target}'


def _attrs_without_type(attrs: str) -> str:
    return re.sub(r'\s+t="[^"]*"', '', attrs)


def _number_text(value: int | float) -> str:
    return str(value) if isinstance(value, int) else repr(value)


def _value_cell(attrs: str, value: NewValue) -> str:
    attrs = _attrs_without_type(attrs)
    if value is None:
        return f'<c{attrs}/>'
    if isinstance(value, bool):
        return f'<c{attrs} t="b"><v>{int(value)}</v></c>'
    if isinstance(value, str):
        space = ' xml:space="preserve"' if value != value.strip() else ''
        return f'<c{attrs} t="inlineStr"><is><t{space}>{escape(value)}</t></is></c>'
    return f'<c{attrs}><v>{_number_text(value)}</v></c>'


def _result_cell(attrs: str, inner: str, value) -> str:
    """Same formula, new stored result."""
    attrs = _attrs_without_type(attrs)
    f = re.search(r'<f\b.*?(?:/>|</f>)', inner, re.S).group()
    if isinstance(value, bool):
        return f'<c{attrs} t="b">{f}<v>{int(value)}</v></c>'
    if isinstance(value, str) or value is None:
        return f'<c{attrs} t="str">{f}<v>{escape(value or "")}</v></c>'
    return f'<c{attrs}>{f}<v>{_number_text(value)}</v></c>'


def _patch_row(row_xml: str, row: int, cells: dict[int, str]) -> str:
    """Replace or insert the given cells ({column: xml}) in one <row>."""
    m = re.match(r'(<row\b[^>]*?)(/?)>', row_xml)
    head, empty = m.group(1), m.group(2)
    body = '' if empty else row_xml[m.end() : -len('</row>')]

    def column_of(match: re.Match) -> int:
        return column_index_from_string(_REF_RE.search(match.group(1)).group(1))

    existing = {column_of(c): c for c in _CELL_RE.finditer(body)}
    for column, xml in sorted(cells.items()):
        if column in existing:
            old = existing[column]
            body = body[: old.start()] + xml + body[old.end() :]
            existing = {column_of(c): c for c in _CELL_RE.finditer(body)}
            continue
        after = [c for col, c in sorted(existing.items()) if col > column]
        at = after[0].start() if after else len(body)
        body = body[:at] + xml + body[at:]
        existing = {column_of(c): c for c in _CELL_RE.finditer(body)}

    spans = re.search(r'\bspans="(\d+):(\d+)"', head)
    if spans:
        low = min(int(spans.group(1)), *cells)
        high = max(int(spans.group(2)), *cells)
        head = head[: spans.start()] + f'spans="{low}:{high}"' + head[spans.end() :]
    return f'{head}>{body}</row>'


def _patch_sheet(xml: str, new: dict, results: dict) -> str:
    by_row: dict[int, dict[int, str]] = {}
    for (r, c), value in new.items():
        old = _find_cell(xml, r, c)
        attrs = old.group(1) if old else f' r="{_coord(r, c)}"'
        by_row.setdefault(r, {})[c] = _value_cell(attrs, value)
    for (r, c), value in results.items():
        old = _find_cell(xml, r, c)
        by_row.setdefault(r, {})[c] = _result_cell(old.group(1), old.group(2), value)

    for r, cells in by_row.items():
        row = re.search(rf'<row\b[^>]*\br="{r}"[^>]*?(?:/>|>.*?</row>)', xml, re.S)
        if row is None:
            raise WriteRefused(f'row {r} does not exist in the sheet')
        xml = xml[: row.start()] + _patch_row(row.group(), r, cells) + xml[row.end() :]
    return xml


def _find_cell(xml: str, row: int, column: int) -> re.Match | None:
    return re.search(
        rf'<c\b([^>]*?\br="{_coord(row, column)}"[^>]*?)(?:/>|>(.*?)</c>)', xml, re.S
    )


def _patch(xlsx: bytes, new: dict, results: dict) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as src, zipfile.ZipFile(out, 'w') as dst:
        sheet = _first_sheet_part(src)
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == sheet:
                data = _patch_sheet(data.decode('utf-8'), new, results).encode('utf-8')
            dst.writestr(item, data)
    return out.getvalue()


# --- verify ------------------------------------------------------------------


def _key(ch: CellChange) -> tuple:
    return (ch.sheet, ch.row, ch.column, ch.field)


def verify(before: bytes, after: bytes, expected: list[CellChange]) -> None:
    """Raise unless `after` is `before` plus exactly `expected`: every
    part except the edited sheet byte-identical, every cell change listed."""
    books = (
        load_workbook(io.BytesIO(before)),
        load_workbook(io.BytesIO(before), data_only=True),
    )
    _verify(before, after, expected, books)


def _verify(
    before: bytes, after: bytes, expected: list[CellChange], before_books
) -> None:
    with (
        zipfile.ZipFile(io.BytesIO(before)) as a,
        zipfile.ZipFile(io.BytesIO(after)) as b,
    ):
        if a.namelist() != b.namelist():
            raise VerifyFailed('the file parts changed')
        sheet = _first_sheet_part(a)
        changed = [n for n in a.namelist() if n != sheet and a.read(n) != b.read(n)]
        if changed:
            raise VerifyFailed(f'parts changed that must not: {changed}')

    after_books = (
        load_workbook(io.BytesIO(after)),
        load_workbook(io.BytesIO(after), data_only=True),
    )
    actual = sorted(changes_between(before_books, after_books), key=_key)
    wanted = sorted(expected, key=_key)
    if actual != wanted:
        extra = [ch for ch in actual if ch not in wanted]
        missing = [ch for ch in wanted if ch not in actual]
        raise VerifyFailed(f'unexpected changes {extra}, missing changes {missing}')
