"""One entry per column and per sample: the only vocabulary the LLM may use.

Everything is read from the file: no column, unit or type is assumed.
"""

import datetime
import re
from dataclasses import dataclass
from typing import Literal

from openpyxl.worksheet.worksheet import Worksheet

from sand.sheet_edit.formulas import Unsupported, references
from sand.sheet_edit.workbook import (
    NOMAD_ID_HEADER,
    SAMPLE_HEADER,
    SUBBATCH_HEADER,
    Column,
    SheetStructure,
    first_sheet,
    info_column,
    sample_label,
    sheet_structure,
)
from sand.sheet_headers import INDEXED_GROUPS, STRIP_RE, classify, norm_key, step_key

ValueType = Literal[
    'empty', 'int', 'float', 'text', 'bool', 'date', 'yyyymmdd', 'mixed'
]

_UNIT_RE = re.compile(r'\[([^\]]+)\]')
_NUMBER_RE = re.compile(r'(?<![\w.])\d+(?![\w.])')
_YYYYMMDD_DIGITS = 8
_LABEL_RE = re.compile(r'(.*?)(\d+)')


@dataclass(frozen=True)
class Family:
    """Columns that describe one thing. 'Solvent 1 name' and 'Solvent 1
    volume' are (solvents, 1); a step's other numbered columns share only
    the number (Seq-Evaporation 'Material name 1', 'Thickness 1')."""

    group: str | None  # 'solvents', 'rotation_steps', ...; None: number only
    index: int


@dataclass(frozen=True)
class ValueGroup:
    """E.g. ValueGroup(value='15min', rows=(3, 4, 5))."""

    value: object  # as stored (a formula's stored result); None: empty
    rows: tuple[int, ...]  # sample rows holding it


@dataclass(frozen=True)
class ColumnInfo:
    id: str  # 'C033': from the column's place in the sheet, which edits never change
    column: Column
    family: Family | None  # 'Solvent 1 volume [uL]' -> Family('solvents', 1)
    unit: str | None  # from '[...]' in the header
    type: ValueType  # of the values in the sample rows
    formula: bool  # a sample cell holds a formula: not writable
    inputs: tuple[str, ...] | None  # ids the formulas read; None: unknown
    groups: tuple[ValueGroup, ...]  # in order of first appearance


@dataclass(frozen=True)
class SampleInfo:
    row: int
    label: str | None  # the Sample value as users say it: '23', never the row
    subbatch: str | None
    nomad_id: str | None  # Excel's stored result of the Nomad ID formula


@dataclass(frozen=True)
class Catalog:
    structure: SheetStructure
    columns: tuple[ColumnInfo, ...]
    samples: tuple[SampleInfo, ...]

    @property
    def needs_subbatch(self) -> bool:
        """Labels repeat (each subbatch restarts at 1): 'sample 1' alone
        names more than one row."""
        labels = [s.label for s in self.samples if s.label is not None]
        return len(labels) != len(set(labels))


def column_id(index: int) -> str:
    return f'C{index:03d}'


def build(xlsx: bytes) -> Catalog:
    ws = first_sheet(xlsx)
    results = first_sheet(xlsx, data_only=True)
    structure = sheet_structure(ws)
    rows = [s.row for s in structure.samples]
    headed = {col.index for col in structure.columns}
    columns = []
    for col in structure.columns:
        values = [results.cell(r, col.index).value for r in rows]
        columns.append(
            ColumnInfo(
                id=column_id(col.index),
                column=col,
                family=_family(col),
                unit=_unit(col.header),
                type=value_type(col.header, values),
                formula=any(ws.cell(r, col.index).data_type == 'f' for r in rows),
                inputs=_inputs(ws, col.index, rows, headed),
                groups=_groups(rows, values),
            )
        )
    return Catalog(structure, tuple(columns), _samples(results, structure))


def _samples(results: Worksheet, structure: SheetStructure) -> tuple[SampleInfo, ...]:
    def value(row: int, header: str) -> str | None:
        col = info_column(structure.columns, header)
        return sample_label(results.cell(row, col).value) if col else None

    return tuple(
        SampleInfo(
            s.row,
            value(s.row, SAMPLE_HEADER),
            value(s.row, SUBBATCH_HEADER),
            value(s.row, NOMAD_ID_HEADER),
        )
        for s in structure.samples
    )


def _groups(rows: list[int], values: list) -> tuple[ValueGroup, ...]:
    found: dict[tuple, tuple[object, list[int]]] = {}
    for row, v in zip(rows, values):
        _, group_rows = found.setdefault(_same(v), (_stored(v), []))
        group_rows.append(row)
    return tuple(ValueGroup(v, tuple(rs)) for v, rs in found.values())


def _stored(v):
    return None if _kind(v) == 'empty' else v


def _same(v) -> tuple:
    """Values Excel holds as equal: 3000 and 3000.0 are one number, but
    TRUE is not 1 and the number 3000 is not the text '3000'."""
    kind = _kind(v)
    if kind in ('int', 'float'):
        return ('number', float(v))
    return (kind, _stored(v))


def sample_ranges(labels: list[str]) -> str:
    """'1–3, 7, 9–12' or 's_1–s_3': consecutive numbers joined, in the
    given order. A missing sample always breaks a range."""
    runs: list[list[str]] = []
    for label in labels:
        if runs and _follows(runs[-1][-1], label):
            runs[-1].append(label)
        else:
            runs.append([label])
    return ', '.join(r[0] if len(r) == 1 else f'{r[0]}–{r[-1]}' for r in runs)


def _follows(previous: str, label: str) -> bool:
    a, b = _numbered(previous), _numbered(label)
    return a is not None and b is not None and b == (a[0], a[1] + 1)


def _numbered(label: str) -> tuple[str, int] | None:
    """'7' -> ('', 7), 's_7' -> ('s_', 7). '07', '1.5' and 'A' are not
    numbered: they are kept as written, never joined."""
    m = _LABEL_RE.fullmatch(label)
    if not m or m.group(2) != str(int(m.group(2))) or m.group(1).endswith(('.', '-')):
        return None
    return m.group(1), int(m.group(2))


def _family(col: Column) -> Family | None:
    spec = classify(col.header, INDEXED_GROUPS.get(step_key(col.section.step_type)))
    if spec['kind'] == 'list':
        return Family(spec['list'], spec['idx'])
    numbers = _NUMBER_RE.findall(STRIP_RE.sub('', col.header))
    return Family(None, int(numbers[0])) if len(numbers) == 1 else None


def _unit(header: str) -> str | None:
    m = _UNIT_RE.search(header)
    return m.group(1).strip() if m else None


def value_type(header: str, values: list) -> ValueType:
    kinds = {_kind(v) for v in values} - {'empty'}
    if kinds == {'int'} and _is_date_column(header, values):
        return 'yyyymmdd'
    if kinds == {'int', 'float'}:
        return 'float'
    if len(kinds) > 1:
        return 'mixed'
    return kinds.pop() if kinds else 'empty'


def _kind(v) -> ValueType:
    if v is None or (isinstance(v, str) and not v.strip()):
        return 'empty'
    if isinstance(v, bool):
        return 'bool'
    if isinstance(v, int) or (isinstance(v, float) and v.is_integer()):
        return 'int'
    if isinstance(v, float):
        return 'float'
    if isinstance(v, datetime.date | datetime.time | datetime.timedelta):
        return 'date'
    return 'text'


def _is_date_column(header: str, values: list) -> bool:
    """KIT writes dates as the number 20260914. Both must hold: the header
    says date, and every value is a real calendar day."""
    if 'date' not in norm_key(header).split('_'):
        return False
    return all(_is_day(str(int(v))) for v in values if _kind(v) == 'int')


def _is_day(digits: str) -> bool:
    if len(digits) != _YYYYMMDD_DIGITS:
        return False
    try:
        datetime.datetime.strptime(digits, '%Y%m%d')
    except ValueError:
        return False
    return True


def _inputs(
    ws: Worksheet, column: int, rows: list[int], headed: set[int]
) -> tuple[str, ...] | None:
    read: set[int] = set()
    for r in rows:
        cell = ws.cell(r, column)
        if cell.data_type != 'f':
            continue
        if not isinstance(cell.value, str):  # array or data-table formula
            return None
        try:
            read |= {ref.column for ref in references(cell.value)}
        except Unsupported:
            return None
    if not read <= headed:  # reads a cell outside the catalogued columns
        return None
    return tuple(column_id(c) for c in sorted(read))
