"""Read the layout of a hysprint-style sheet from the file itself.

There is no fixed template: row 1 names the step groups, row 2 the
columns, rows 3+ are samples. Groups are found the way NOMAD's batch
parser finds them (pandas `header=[0, 1]`): a row-1 label opens a group
that runs until the next label, so merged cells are not required.
"""

import io
import re
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

GROUP_ROW = 1
HEADER_ROW = 2
FIRST_SAMPLE_ROW = 3
EXPERIMENT_INFO = 'Experiment Info'
SAMPLE_HEADER = 'Sample'
SUBBATCH_HEADER = 'Subbatch'
NOMAD_ID_HEADER = 'Nomad ID'

_POSITION_RE = re.compile(r'^(\d+)\s*:\s*(.+)$')


@dataclass(frozen=True)
class Section:
    label: str  # as written in row 1, e.g. '2: Spin Coating'
    step_type: str  # 'Spin Coating'; 'Experiment Info' for the info group
    position: int | None  # the 'N:' prefix; None for Experiment Info
    first_col: int  # 1-based, inclusive
    last_col: int


@dataclass(frozen=True)
class Column:
    index: int  # 1-based: column A is 1
    header: str  # the column name in row 2, e.g. 'Rotation speed [rpm]'
    section: Section


@dataclass(frozen=True)
class SampleRow:
    row: int  # 1-based sheet row
    sample: str | None  # this row's value in the Sample column, e.g. '7'


@dataclass(frozen=True)
class SheetStructure:
    sheet_name: str
    sections: tuple[Section, ...]
    columns: tuple[Column, ...]
    samples: tuple[SampleRow, ...]


def _text(value) -> str:
    return '' if value is None else str(value).strip()


def _sample(value) -> str | None:
    """Sample names as users say them: 1.0 and 1 are both '1'."""
    if value is None or _text(value) == '':
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _section(label: str, first_col: int, last_col: int) -> Section:
    m = _POSITION_RE.match(label)
    if m:
        return Section(label, m.group(2).strip(), int(m.group(1)), first_col, last_col)
    return Section(label, label, None, first_col, last_col)


def _info_column(columns: tuple[Column, ...], header: str) -> int | None:
    return next(
        (
            col.index
            for col in columns
            if col.section.step_type == EXPERIMENT_INFO and col.header == header
        ),
        None,
    )


def first_sheet(xlsx: bytes, data_only: bool = False) -> Worksheet:
    """data_only: formula cells hold Excel's stored result, not the formula."""
    return load_workbook(io.BytesIO(xlsx), data_only=data_only).worksheets[0]


def read_structure(xlsx: bytes) -> SheetStructure:
    """Sections, columns and sample rows of the first sheet (the one the
    batch parser reads)."""
    return sheet_structure(first_sheet(xlsx))


def sheet_structure(ws: Worksheet) -> SheetStructure:
    width = max(
        (c for c in range(1, ws.max_column + 1) if _text(ws.cell(HEADER_ROW, c).value)),
        default=0,
    )

    starts = [c for c in range(1, width + 1) if _text(ws.cell(GROUP_ROW, c).value)]
    sections = tuple(
        _section(_text(ws.cell(GROUP_ROW, a).value), a, b - 1)
        for a, b in zip(starts, [*starts[1:], width + 1])
    )

    columns = tuple(
        Column(c, _text(ws.cell(HEADER_ROW, c).value), sec)
        for sec in sections
        for c in range(sec.first_col, sec.last_col + 1)
        if _text(ws.cell(HEADER_ROW, c).value)
    )

    sample_col = _info_column(columns, SAMPLE_HEADER)
    samples = tuple(
        SampleRow(
            r, _sample(ws.cell(r, sample_col).value) if sample_col is not None else None
        )
        for r in range(FIRST_SAMPLE_ROW, ws.max_row + 1)
        if any(_text(ws.cell(r, c).value) for c in range(1, width + 1))
    )
    return SheetStructure(ws.title, sections, columns, samples)


# --- ingest checks ----------------------------------------------------------


def ingest_problems(xlsx: bytes) -> list[str]:
    """Why this sheet cannot be edited safely; empty means it can.

    Every check guards addressing: a cell must have exactly one
    (step, column, sample) address, or an edit could land in the wrong
    place.
    """
    try:
        ws = first_sheet(xlsx)
    except Exception as exc:  # any openpyxl/zip failure means "not an xlsx"
        return [f'not a readable .xlsx file ({exc})']
    s = sheet_structure(ws)
    if not s.columns:
        return ['row 2 has no column headers']

    problems = [
        *_group_problems(ws, s),
        *_column_problems(ws, s),
        *_sample_problems(ws, s),
    ]
    problems += [
        f'merged cells {rng.coord} below row 1: cells must be single to be edited'
        for rng in ws.merged_cells.ranges
        if rng.max_row > GROUP_ROW
    ]
    return problems


def _group_problems(ws: Worksheet, s: SheetStructure) -> list[str]:
    problems = []
    first = s.sections[0] if s.sections else None
    if first is None or first.first_col != 1 or first.label != EXPERIMENT_INFO:
        problems.append(
            f"cell A1 must be '{EXPERIMENT_INFO}' "
            f'(found {_text(ws.cell(GROUP_ROW, 1).value)!r})'
        )
    seen: dict[int, str] = {}
    for sec in s.sections[1:]:
        where = get_column_letter(sec.first_col)
        if sec.position is None:
            problems.append(
                f"step label {sec.label!r} (cell {where}1) has no 'N:' position"
            )
        elif sec.position in seen:
            problems.append(
                f'step position {sec.position} is used twice '
                f'(cells {seen[sec.position]}1 and {where}1)'
            )
        else:
            seen[sec.position] = where
    return problems


def _column_problems(ws: Worksheet, s: SheetStructure) -> list[str]:
    problems = []
    headed = {col.index for col in s.columns}
    for c in range(1, ws.max_column + 1):
        if c not in headed and any(
            _text(ws.cell(r, c).value) for r in range(FIRST_SAMPLE_ROW, ws.max_row + 1)
        ):
            problems.append(
                f'column {get_column_letter(c)} has values but no header in row 2'
            )
    for sec in s.sections:
        by_header: dict[str, list[str]] = {}
        for col in s.columns:
            if col.section == sec:
                by_header.setdefault(col.header, []).append(
                    get_column_letter(col.index)
                )
        problems += [
            f'{header!r} appears twice in {sec.label!r} (columns {", ".join(cols)})'
            for header, cols in by_header.items()
            if len(cols) > 1
        ]
    return problems


def _sample_problems(ws: Worksheet, s: SheetStructure) -> list[str]:
    problems = [
        f"'{EXPERIMENT_INFO}' has no {header!r} column"
        for header in (NOMAD_ID_HEADER, SAMPLE_HEADER)
        if _info_column(s.columns, header) is None
    ]
    if not s.samples:
        return [*problems, 'no sample rows below the headers']

    rows = [r.row for r in s.samples]
    problems += [
        f'row {r} is empty between sample rows'
        for r in range(rows[0], rows[-1] + 1)
        if r not in set(rows)
    ]
    if _info_column(s.columns, SAMPLE_HEADER) is None:
        return problems

    # Samples may restart at 1 in each subbatch: the address is then
    # (subbatch, sample), and a bare "sample 1" is resolved later.
    subbatch_col = _info_column(s.columns, SUBBATCH_HEADER)
    seen: dict[tuple, int] = {}
    for r in s.samples:
        if r.sample is None:
            problems.append(f'row {r.row} has no {SAMPLE_HEADER!r} value')
            continue
        subbatch = _sample(ws.cell(r.row, subbatch_col).value) if subbatch_col else None
        key = (subbatch, r.sample)
        if key in seen:
            name = f'sample {r.sample}' + (
                f' of subbatch {subbatch}' if subbatch is not None else ''
            )
            problems.append(f'{name} appears twice (rows {seen[key]} and {r.row})')
        else:
            seen[key] = r.row
    return problems


# --- round trip ---------------------------------------------------------------

# Package parts that hold no lab data: Excel's calculation-order cache,
# the shared-string table (cell text is compared cell by cell instead),
# print setup, and relationship files (their targets are compared).
_NON_DATA_PARTS = re.compile(
    r'^xl/calcChain\.xml$|^xl/sharedStrings\.xml$|^xl/printerSettings/|\.rels$'
)


@dataclass(frozen=True)
class CellChange:
    sheet: str
    row: int  # 1-based, like SampleRow.row
    column: int  # 1-based, like Column.index
    field: str  # 'value' | 'type' | 'number format' | 'computed value'
    before: object
    after: object


def cell_changes(before: bytes, after: bytes) -> list[CellChange]:
    """Every cell difference on sheets present in both files.

    'computed value' is the result Excel cached for a formula. openpyxl
    cannot compute formulas, and NOMAD's parser reads only these cached
    results, so losing one loses data even though the formula is intact.
    """
    return changes_between(
        (
            load_workbook(io.BytesIO(before)),
            load_workbook(io.BytesIO(before), data_only=True),
        ),
        (
            load_workbook(io.BytesIO(after)),
            load_workbook(io.BytesIO(after), data_only=True),
        ),
    )


def _comparable(value):
    """Array and data-table formulas are objects equal only to themselves;
    compare them by what they hold."""
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if hasattr(value, '__dict__'):
        return (type(value).__name__, sorted(vars(value).items()))
    return value


def changes_between(before: tuple, after: tuple) -> list[CellChange]:
    """cell_changes on already loaded (formulas, data_only) workbook pairs."""
    (a, a_vals), (b, b_vals) = before, after
    changes = []
    for name in a.sheetnames:
        if name not in b.sheetnames:
            continue
        wa, wb = a[name], b[name]
        for r in range(1, max(wa.max_row, wb.max_row) + 1):
            for c in range(1, max(wa.max_column, wb.max_column) + 1):
                ca, cb = wa.cell(r, c), wb.cell(r, c)
                for field, x, y in (
                    ('value', _comparable(ca.value), _comparable(cb.value)),
                    ('type', ca.data_type, cb.data_type),
                    ('number format', ca.number_format, cb.number_format),
                ):
                    if x != y:
                        changes.append(CellChange(name, r, c, field, x, y))
                if ca.data_type == 'f':
                    x = a_vals[name].cell(r, c).value
                    y = b_vals[name].cell(r, c).value
                    if x != y:
                        changes.append(CellChange(name, r, c, 'computed value', x, y))
    return changes


# TODO: not sure we need it; if edits patch only the target cells (PR 4),
# "every other part byte-identical" replaces this check.
def _data_parts(xlsx: bytes) -> set[str]:
    names = zipfile.ZipFile(io.BytesIO(xlsx)).namelist()
    return {n for n in names if not _NON_DATA_PARTS.search(n)}


def workbook_problems(before: bytes, after: bytes) -> list[str]:
    """What `after` lost or changed against `before`, as user messages."""
    a, b = load_workbook(io.BytesIO(before)), load_workbook(io.BytesIO(after))
    problems = []
    if a.sheetnames != b.sheetnames:
        problems.append(f'sheets changed: {a.sheetnames} -> {b.sheetnames}')
    for name in set(a.sheetnames) & set(b.sheetnames):
        ma = sorted(str(m) for m in a[name].merged_cells.ranges)
        mb = sorted(str(m) for m in b[name].merged_cells.ranges)
        if ma != mb:
            problems.append(f'{name!r}: merged cells changed')
    problems += [
        f'file part {part} lost'
        for part in sorted(_data_parts(before) - _data_parts(after))
    ]

    # one message per column for lost formula results: the usual case is
    # a whole formula column, not single cells
    lost: dict[tuple[str, int], list[int]] = {}
    for ch in cell_changes(before, after):
        if ch.field == 'computed value' and ch.after is None:
            lost.setdefault((ch.sheet, ch.column), []).append(ch.row)
        else:
            cell = f'{get_column_letter(ch.column)}{ch.row}'
            problems.append(
                f'{ch.sheet!r}!{cell}: {ch.field} {ch.before!r} -> {ch.after!r}'
            )
    for (sheet, column), rows in lost.items():
        col = get_column_letter(column)
        problems.append(
            f'{sheet!r} column {col}: formula results lost on {len(rows)} rows '
            f'({col}{rows[0]}…{col}{rows[-1]})'
        )
    return problems


def roundtrip_problems(xlsx: bytes) -> list[str]:
    """What opening and saving this file with openpyxl, changing nothing,
    would lose. Every later edit is saved the same way."""
    buffer = io.BytesIO()
    load_workbook(io.BytesIO(xlsx)).save(buffer)
    return workbook_problems(xlsx, buffer.getvalue())
