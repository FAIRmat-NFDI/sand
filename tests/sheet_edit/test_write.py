import datetime
import io
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from synthetic_sheet import letter, with_formula_results, with_part

from sand.sheet_edit.workbook import CellChange
from sand.sheet_edit.write import (
    VerifyFailed,
    Write,
    WriteRefused,
    verify,
    write_cells,
)

SHEET = 'Experiment Data'
INFO = 'Experiment Info'
SPIN2 = '2: Spin Coating'
CSS = '5: Close Space Sublimation'


def col(section: str, header: str) -> int:
    return column_index_from_string(letter(section, header))


SPEED = col(SPIN2, 'Rotation speed [rpm]')
SOLVENT = col(SPIN2, 'Solvent 1 name')
CONCENTRATION = col(SPIN2, 'Solute 1 Concentration [mM]')  # empty everywhere
BATCH = col(INFO, 'Batch')
VARIATION = col(INFO, 'Variation')
NOMAD_ID = col(INFO, 'Nomad ID')
DEPOSITION = col(CSS, 'Deposition Time [s]')
OLD_SPEED, NEW_SPEED = 3000, 4000


def cell(xlsx: bytes, row: int, column: int, data_only=False):
    ws = load_workbook(io.BytesIO(xlsx), data_only=data_only).worksheets[0]
    return ws.cell(row, column)


def parts(xlsx: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.fixture
def with_results(synthetic_xlsx) -> bytes:
    """The synthetic sheet with Excel-stored formula results on row 3."""
    return with_formula_results(
        synthetic_xlsx,
        {
            f'{letter(INFO, "Nomad ID")}3': 'KIT_Synth_20250101_B1_0_1',
            f'{letter(CSS, "Deposition Time [s]")}3': '600',
        },
    )


# --- plain values -----------------------------------------------------------


def test_number_edit_changes_only_that_cell(synthetic_xlsx):
    result = write_cells(synthetic_xlsx, [Write(3, SPEED, NEW_SPEED)])
    assert result.changes == [
        CellChange(SHEET, 3, SPEED, 'value', OLD_SPEED, NEW_SPEED)
    ]
    assert cell(result.xlsx, 3, SPEED).value == NEW_SPEED
    before, after = parts(synthetic_xlsx), parts(result.xlsx)
    assert [n for n in before if before[n] != after[n]] == ['xl/worksheets/sheet1.xml']


@pytest.mark.parametrize(
    ('said', 'stored'), [('4000', 4000), (' 7.5 ', 7.5), (2e3, 2000)]
)
def test_numbers_are_written_as_numbers(synthetic_xlsx, said, stored):
    after = write_cells(synthetic_xlsx, [Write(3, SPEED, said)]).xlsx
    value = cell(after, 3, SPEED).value
    assert value == stored and type(value) is type(stored)


def test_text_in_a_number_column_is_refused(synthetic_xlsx):
    with pytest.raises(WriteRefused, match='holds numbers'):
        write_cells(synthetic_xlsx, [Write(3, SPEED, 'fast')])


def test_text_column_stays_text_even_for_a_number(synthetic_xlsx):
    after = write_cells(
        synthetic_xlsx, [Write(3, SOLVENT, 'DMF'), Write(4, SOLVENT, 2)]
    ).xlsx
    assert cell(after, 3, SOLVENT).value == 'DMF'
    assert cell(after, 4, SOLVENT).value == '2'
    assert cell(after, 4, SOLVENT).data_type == 's'


def test_text_keeps_spaces_and_markup_characters(synthetic_xlsx):
    after = write_cells(synthetic_xlsx, [Write(3, SOLVENT, ' <A&B> ')]).xlsx
    assert cell(after, 3, SOLVENT).value == ' <A&B> '


def test_empty_cell_gets_a_new_cell_in_column_order(synthetic_xlsx):
    result = write_cells(synthetic_xlsx, [Write(5, CONCENTRATION, '5.5')])
    # empty column: a number said is stored as a number
    assert result.changes == [CellChange(SHEET, 5, CONCENTRATION, 'value', None, 5.5)]
    ws = load_workbook(io.BytesIO(result.xlsx)).worksheets[0]
    assert ws.cell(5, CONCENTRATION).value == 5.5  # noqa: PLR2004
    assert ws.cell(5, CONCENTRATION - 1).value == 'SAM'  # neighbours intact
    assert ws.cell(5, SPEED).value == OLD_SPEED


def test_clearing_a_cell(synthetic_xlsx):
    result = write_cells(synthetic_xlsx, [Write(3, SOLVENT, None)])
    assert result.changes == [
        CellChange(SHEET, 3, SOLVENT, 'value', 'Ethanol', None),
        CellChange(SHEET, 3, SOLVENT, 'type', 's', 'n'),
    ]


def test_writing_the_same_value_changes_nothing(synthetic_xlsx):
    result = write_cells(synthetic_xlsx, [Write(3, SPEED, '3000')])
    assert result.changes == []
    assert result.xlsx == synthetic_xlsx


@pytest.mark.parametrize(
    ('writes', 'message'),
    [
        ([Write(3, NOMAD_ID, 'x')], 'calculated by a formula'),
        ([Write(2, SPEED, 'x')], 'headers are not edited'),
        ([Write(3, SPEED, 1), Write(3, SPEED, 2)], 'written twice'),
    ],
)
def test_refused_writes(synthetic_xlsx, writes, message):
    with pytest.raises(WriteRefused, match=message):
        write_cells(synthetic_xlsx, writes)


# --- formula results ---------------------------------------------------------


def test_formula_input_edit_recomputes_the_result(with_results):
    result = write_cells(
        with_results, [Write(3, BATCH, 'B2'), Write(3, VARIATION, '30min')]
    )
    assert sorted(result.changes, key=lambda c: c.column) == [
        CellChange(SHEET, 3, BATCH, 'value', 'B1', 'B2'),
        CellChange(
            SHEET,
            3,
            NOMAD_ID,
            'computed value',
            'KIT_Synth_20250101_B1_0_1',
            'KIT_Synth_20250101_B2_0_1',
        ),
        CellChange(SHEET, 3, VARIATION, 'value', '10min', '30min'),
        CellChange(SHEET, 3, DEPOSITION, 'computed value', 600, 1800),
    ]
    stored = load_workbook(io.BytesIO(result.xlsx), data_only=True).worksheets[0]
    assert stored.cell(3, NOMAD_ID).value == 'KIT_Synth_20250101_B2_0_1'
    assert stored.cell(3, DEPOSITION).value == 1800  # noqa: PLR2004
    # the formula itself is untouched
    assert (
        cell(result.xlsx, 3, DEPOSITION).value
        == cell(with_results, 3, DEPOSITION).value
    )


def test_formula_without_a_stored_result_is_not_trusted(synthetic_xlsx):
    # files saved by openpyxl carry no results to check the calculator against
    with pytest.raises(WriteRefused, match='no stored result'):
        write_cells(synthetic_xlsx, [Write(3, BATCH, 'B2')])


def test_calculator_must_reproduce_excels_result(synthetic_xlsx):
    wrong = with_formula_results(
        synthetic_xlsx, {f'{letter(INFO, "Nomad ID")}3': 'KIT_something_else'}
    )
    with pytest.raises(WriteRefused, match="Excel stored 'KIT_something_else'"):
        write_cells(wrong, [Write(3, BATCH, 'B2')])


def test_edit_that_makes_a_formula_error_is_refused(with_results):
    with pytest.raises(WriteRefused, match='would become an Excel error'):
        write_cells(with_results, [Write(3, VARIATION, 'xx')])


def _small(formulas: dict[str, str], results: dict[str, str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(['Experiment Info', None, None, None])
    ws.append(['Sample', 'x', 'f1', 'f2'])
    ws.append([1, 2, None, None])
    ws.append([2, 3, None, None])
    for coord, text in formulas.items():
        ws[coord] = text
    buffer = io.BytesIO()
    wb.save(buffer)
    return with_formula_results(buffer.getvalue(), results)


def test_chained_formulas_are_recomputed_in_order():
    xlsx = _small({'C3': '=B3*2', 'D3': '=C3+1'}, {'C3': '4', 'D3': '5'})
    result = write_cells(xlsx, [Write(3, 2, 10)])
    assert [
        (c.column, c.after) for c in result.changes if c.field == 'computed value'
    ] == [
        (3, 20),
        (4, 21),
    ]


def test_unsupported_formula_blocks_only_edits_it_reads():
    xlsx = _small({'C3': '=SUM(B3:B4)'}, {'C3': '5'})
    with pytest.raises(WriteRefused, match='formula sand cannot compute'):
        write_cells(xlsx, [Write(4, 2, 7)])
    assert write_cells(xlsx, [Write(4, 1, 9)]).changes  # Sample is not read


# --- verify -----------------------------------------------------------------


def test_verify_rejects_an_unlisted_change(synthetic_xlsx):
    result = write_cells(synthetic_xlsx, [Write(3, SPEED, 4000), Write(4, SPEED, 4000)])
    with pytest.raises(VerifyFailed, match='unexpected changes'):
        verify(synthetic_xlsx, result.xlsx, result.changes[:1])


def test_verify_rejects_a_changed_part(synthetic_xlsx):
    result = write_cells(synthetic_xlsx, [Write(3, SPEED, 4000)])
    with pytest.raises(VerifyFailed, match='file parts changed'):
        verify(synthetic_xlsx, with_part(result.xlsx, 'xl/extra.xml'), result.changes)


# --- the real KIT sheet --------------------------------------------------------


def test_real_kit_sheet_edits_reach_nomads_reader(real_kit_xlsx):
    batch, variation, speed, conc = 3, 7, 33, 30  # C, G, AG, AD
    result = write_cells(
        real_kit_xlsx,
        [
            Write(3, batch, '276nmPbI2'),
            Write(4, variation, '50min'),
            Write(3, speed, '4000'),
            Write(5, conc, 5),  # an empty cell
        ],
    )
    df = pd.read_excel(io.BytesIO(result.xlsx), header=[0, 1])
    ids = df['Experiment Info']['Nomad ID'].to_list()
    assert ids[:2] == [
        'KIT_JoDa_20260914_276nmPbI2_0_1',
        'KIT_JoDa_20260914_275nmPbI2_0_2',
    ]
    assert len([i for i in ids if isinstance(i, str)]) == 42  # noqa: PLR2004
    times = df['5: Close Space Sublimation']['Deposition Time [s]'].to_list()
    assert times[:3] == [900, 3000, 900]


# --- review fixes: cases sand cannot handle exactly are refused --------------


def _book(rows: list[list], results: dict[str, str] | None = None, setup=None) -> bytes:
    """A small sheet: two header rows, then `rows` from row 3."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet'
    ws.append(['Experiment Info', None, None, None])
    ws.append(['Sample', 'x', 'f1', 'f2'])
    for row in rows:
        ws.append(row)
    if setup:
        setup(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return with_formula_results(buffer.getvalue(), results or {})


def test_defined_name_makes_every_edit_feed_the_formula():
    def name(wb):
        wb.defined_names['Speed'] = DefinedName('Speed', attr_text='Sheet!$B$3')

    xlsx = _book([[1, 2, '=Speed*2']], {'C3': '4'}, name)
    for column in (1, 2):  # sand cannot tell what Speed points at
        with pytest.raises(WriteRefused, match='formula sand cannot compute'):
            write_cells(xlsx, [Write(3, column, 5)])


def test_indirect_reads_anything():
    xlsx = _book([[1, 2, '=INDIRECT("B3")*2']], {'C3': '4'})
    with pytest.raises(WriteRefused, match='formula sand cannot compute'):
        write_cells(xlsx, [Write(3, 1, 5)])


def test_whole_row_range_is_seen():
    xlsx = _book([[1, 2, '=SUM(3:3)'], [2, 3, None]], {'C3': '3'})
    with pytest.raises(WriteRefused, match='formula sand cannot compute'):
        write_cells(xlsx, [Write(3, 2, 5)])
    assert write_cells(xlsx, [Write(4, 2, 5)]).changes  # row 4 is not read


def test_date_cell_is_not_turned_into_text():
    xlsx = _book([[1, datetime.datetime(2025, 1, 1), None]])
    with pytest.raises(WriteRefused, match='holds a date'):
        write_cells(xlsx, [Write(3, 2, '2025-02-01')])


def test_formula_over_a_date_is_refused_not_crashing():
    xlsx = _book(
        [['a', datetime.datetime(2025, 1, 1), '=A3&"_"&B3']],
        {'C3': 'a_45658'},
    )
    with pytest.raises(WriteRefused, match='date values are not supported'):
        write_cells(xlsx, [Write(3, 1, 'b')])


def test_circular_formulas_are_refused():
    xlsx = _book([[1, 2, '=D3+B3', '=C3*0+1']], {'C3': '3', 'D3': '1'})
    with pytest.raises(WriteRefused, match='circular'):
        write_cells(xlsx, [Write(3, 2, 5)])


@pytest.mark.parametrize(('row', 'column'), [(3, 0), (0, 1), (3, 10**5), (10**6, 2)])
def test_writes_outside_the_sheet_are_refused(row, column):
    xlsx = _book([[1, 2, None]])
    with pytest.raises(WriteRefused, match='outside the sheet'):
        write_cells(xlsx, [Write(row, column, 1)])


def test_empty_formula_result_is_a_valid_stored_result():
    xlsx = _book(
        [[1, 'a', '=IF(B3="x","yes","")'], [2, 'x', '=IF(B4="x","yes","")']],
        {'C3': '', 'C4': 'yes'},
    )
    result = write_cells(xlsx, [Write(3, 2, 'x'), Write(4, 2, 'b')])
    computed = {
        (c.row, c.before, c.after)
        for c in result.changes
        if c.field == 'computed value'
    }
    assert computed == {(3, None, 'yes'), (4, 'yes', None)}
    # still empty -> no change listed
    assert write_cells(xlsx, [Write(3, 2, 'c')]).changes == [
        CellChange('Sheet', 3, 2, 'value', 'a', 'c')
    ]


def test_array_formula_blocks_only_edits_it_reads():
    def array(wb):
        wb.active['D3'] = ArrayFormula('D3', '=SUM(B3:B4)')

    xlsx = _book([[1, 2, None], [2, 3, None]], setup=array)
    assert write_cells(xlsx, [Write(3, 1, 9)]).changes  # Sample is not read
    with pytest.raises(WriteRefused, match='array or table formula'):
        write_cells(xlsx, [Write(3, 2, 9)])


def test_formula_on_another_sheet_pointing_here_is_refused():
    def other(wb):
        wb.create_sheet('Summary')['A1'] = '=Sheet!B3*2'

    xlsx = _book([[1, 2, None]], setup=other)
    with pytest.raises(WriteRefused, match="sheet 'Summary'"):
        write_cells(xlsx, [Write(3, 2, 5)])
