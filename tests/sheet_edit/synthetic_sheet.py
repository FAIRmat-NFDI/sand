"""A synthetic sheet with the KIT lab's layout and invented values.

The real KIT sheet stays private (docs/ is local-only); this one mirrors
what matters for editing: repeated step types, numbered column families,
formula columns, per-sample variation and a gap in the sample labels.
"""

import io
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

SPIN_COATING = [
    'Material name',
    'Layer type',
    'Tool/GB name',
    'Solvent 1 name',
    'Solvent 1 volume [uL]',
    'Solute 1 type',
    'Solute 1 Concentration [mM]',
    'Solution volume [uL]',
    'Rotation speed [rpm]',
    'Rotation time [s]',
    'Annealing time [min]',
    'Annealing temperature [°C]',
    'Notes',
]
SEQ_EVAPORATION = [
    'Material name',
    'Layer type',
    'Material name 1',
    'Thickness 1 [nm]',
    'Rate 1 [angstrom/s]',
    'Material name 2',
    'Thickness 2 [nm]',
    'Rate 2 [angstrom/s]',
    'Notes',
]

SECTIONS = [
    (
        'Experiment Info',
        ['Date', 'Project_Name', 'Batch', 'Subbatch', 'Sample', 'Nomad ID']
        + ['Variation', 'Sample area [cm^2]', 'Notes'],
    ),
    (
        '1: Cleaning UV-Ozone',
        ['Solvent 1', 'Time 1 [s]', 'Temperature 1 [°C]']
        + ['Solvent 2', 'Time 2 [s]', 'Temperature 2 [°C]', 'UV-Ozone Time [s]'],
    ),
    ('2: Spin Coating', SPIN_COATING),
    ('3: Spin Coating', SPIN_COATING),
    ('4: Seq-Evaporation', SEQ_EVAPORATION),
    (
        '5: Close Space Sublimation',
        ['Material name', 'Source material name 1', 'Source material name 2']
        + ['Substrate temperature [°C]', 'Deposition Time [s]', 'Notes'],
    ),
    ('6: Annealing', ['Annealing time [min]', 'Annealing temperature [°C]', 'Notes']),
    ('7: Spin Coating', SPIN_COATING),  # left empty, like the KIT sheet
    ('8: Evaporation', ['Material name', 'Thickness [nm]', 'Rate [angstrom/s]']),
]

# 12 samples, no sample 6; variation in groups of three
SAMPLE_LABELS = [1, 2, 3, 4, 5, 7, 8, 9, 10, 11, 12, 13]
VARIATIONS = ['10min'] * 3 + ['20min'] * 3 + ['10min-hot'] * 3 + ['20min-hot'] * 3

# same value on every sample row, keyed by (section label, header)
CONSTANTS = {
    ('Experiment Info', 'Date'): 20250101,
    ('Experiment Info', 'Project_Name'): 'Synth',
    ('Experiment Info', 'Batch'): 'B1',
    ('Experiment Info', 'Subbatch'): 0,
    ('Experiment Info', 'Sample area [cm^2]'): 0.1,
    ('1: Cleaning UV-Ozone', 'Solvent 1'): 'Acetone',
    ('1: Cleaning UV-Ozone', 'Time 1 [s]'): 600,
    ('1: Cleaning UV-Ozone', 'Temperature 1 [°C]'): 35,
    ('1: Cleaning UV-Ozone', 'Solvent 2'): 'IPA',
    ('1: Cleaning UV-Ozone', 'Time 2 [s]'): 600,
    ('1: Cleaning UV-Ozone', 'Temperature 2 [°C]'): 35,
    ('1: Cleaning UV-Ozone', 'UV-Ozone Time [s]'): 900,
    ('2: Spin Coating', 'Material name'): 'Me-4PACz',
    ('2: Spin Coating', 'Layer type'): 'HTL',
    ('2: Spin Coating', 'Tool/GB name'): 'box1',
    ('2: Spin Coating', 'Solvent 1 name'): 'Ethanol',
    ('2: Spin Coating', 'Solute 1 type'): 'SAM',
    ('2: Spin Coating', 'Solution volume [uL]'): 50,
    ('2: Spin Coating', 'Rotation speed [rpm]'): 3000,
    ('2: Spin Coating', 'Rotation time [s]'): 30,
    ('2: Spin Coating', 'Annealing time [min]'): 10,
    ('2: Spin Coating', 'Annealing temperature [°C]'): 100,
    ('3: Spin Coating', 'Material name'): 'Ethanol wash',
    ('3: Spin Coating', 'Solvent 1 name'): 'Ethanol',
    ('3: Spin Coating', 'Solvent 1 volume [uL]'): 50,
    ('3: Spin Coating', 'Rotation speed [rpm]'): 4000,
    ('3: Spin Coating', 'Rotation time [s]'): 20,
    ('3: Spin Coating', 'Annealing temperature [°C]'): 100,
    ('4: Seq-Evaporation', 'Material name'): 'PbI2',
    ('4: Seq-Evaporation', 'Material name 2'): 'PbI2',
    ('4: Seq-Evaporation', 'Thickness 2 [nm]'): 200,
    ('4: Seq-Evaporation', 'Rate 2 [angstrom/s]'): 1,
    ('5: Close Space Sublimation', 'Material name'): 'MAI+MABr',
    ('5: Close Space Sublimation', 'Source material name 1'): 'MAI',
    ('5: Close Space Sublimation', 'Source material name 2'): 'MABr',
    ('6: Annealing', 'Annealing time [min]'): 10,
    ('6: Annealing', 'Annealing temperature [°C]'): 130,
    ('8: Evaporation', 'Material name'): 'Ag',
    ('8: Evaporation', 'Thickness [nm]'): 100,
}


def build() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Experiment Data'

    col = 1
    where: dict[tuple[str, str], str] = {}  # (section, header) -> column letter
    for label, headers in SECTIONS:
        first = col
        ws.cell(1, first, label)
        for header in headers:
            ws.cell(2, col, header)
            where[label, header] = get_column_letter(col)
            col += 1
        if col - 1 > first:
            ws.merge_cells(
                start_row=1, start_column=first, end_row=1, end_column=col - 1
            )

    info = 'Experiment Info'
    css = '5: Close Space Sublimation'
    for i, (sample, variation) in enumerate(zip(SAMPLE_LABELS, VARIATIONS)):
        r = 3 + i
        for (section, header), value in CONSTANTS.items():
            ws[f'{where[section, header]}{r}'] = value
        ws[f'{where[info, "Sample"]}{r}'] = sample
        ws[f'{where[info, "Variation"]}{r}'] = variation
        ws[f'{where[css, "Substrate temperature [°C]"]}{r}'] = (
            120 if i < len(SAMPLE_LABELS) // 2 else 115
        )
        b, a, c, d, e = (
            where[info, h]
            for h in ('Project_Name', 'Date', 'Batch', 'Subbatch', 'Sample')
        )
        ws[f'{where[info, "Nomad ID"]}{r}'] = (
            f'=CONCATENATE("KIT_",{b}{r},"_",{a}{r},"_",{c}{r},"_",{d}{r},"_",{e}{r})'
        )
        g = where[info, 'Variation']
        ws[f'{where[css, "Deposition Time [s]"]}{r}'] = f'=VALUE(LEFT({g}{r},2))*60'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def letter(section: str, header: str) -> str:
    """Column letter of (section label, header) in the built sheet."""
    col = 1
    for label, headers in SECTIONS:
        for h in headers:
            if (label, h) == (section, header):
                return get_column_letter(col)
            col += 1
    raise KeyError((section, header))


def edited(xlsx: bytes, change) -> bytes:
    """A copy of `xlsx` after `change(worksheet)`."""
    wb = load_workbook(io.BytesIO(xlsx))
    change(wb.worksheets[0])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def with_part(xlsx: bytes, name: str, data: bytes = b'x') -> bytes:
    """A copy with one more package part (e.g. an image or calcChain)."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(xlsx)) as src,
        zipfile.ZipFile(out, 'w') as dst,
    ):
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(name, data)
    return out.getvalue()


def with_formula_results(xlsx: bytes, results: dict[str, str]) -> bytes:
    """A copy where formula cells carry cached results, as Excel saves them.

    openpyxl cannot write these; this patches the first sheet's XML.
    """
    sheet = 'xl/worksheets/sheet1.xml'
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(xlsx)) as src,
        zipfile.ZipFile(out, 'w') as dst,
    ):
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == sheet:
                xml = data.decode()
                for coord, value in results.items():
                    xml, n = re.subn(
                        rf'<c r="{coord}">(<f>.*?</f>)<v></v></c>',
                        rf'<c r="{coord}" t="str">\1<v>{value}</v></c>',
                        xml,
                    )
                    assert n == 1, coord
                data = xml.encode()
            dst.writestr(item, data)
    return out.getvalue()
